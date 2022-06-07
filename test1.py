# -*- coding: utf-8 -*-
"""
@author: drodriguez
"""
#packages for database manipulation
import pyodbc
import pandas as pd
import shutil
import os
import numpy as np
import xlsxwriter
import xlrd
import matplotlib.pyplot as plt



#database manipulation
connection = pyodbc.connect('DRIVER={SQL Server}; SERVER=saaws-sql14; DATABASE=SolveComposite; Trusted_Connection=yes;')

cursor = connection.cursor()
cursor.execute('EXECUTE SolveComposite.dbo.[spCorpCompositeReport] @Sector=\'IG\'')

while 1:
    row = cursor.fetchone()
    if not row:
        break
    print(row)



#Creating prior weekly summaries and backup
#source directory of the files
src_dir = 'D:/Users/drodriguez/OneDrive - Solve Advisors Inc/Documents/Python Scripts/IG & HY current week/'

#destination directory for the files
dest_dir = 'D:/Users/drodriguez/OneDrive - Solve Advisors Inc/Documents/Python Scripts/Prior Weekly Summaries2/'

#getting all the files from the source directory
files = os.listdir(src_dir)
shutil.copytree(src_dir, dest_dir)


#importing workbook load function from openpyx
from openpyxl import load_workbook, Workbook

values = np.array([["20-01-2018",4,9,16,25,36]])
columns=["Date","A","B","C","D","E"]
df = pd.DataFrame(values.transpose(),columns)
path = r'D:\Users\drodriguez\OneDrive - Solve Advisors Inc\Documents\Python Scripts\HY CORP C&G 2022.01.14.xlsx'
df2 = pd.read_excel(path, sheet_name = None)
print(df2)
writer = pd.ExcelWriter(path, engine='openpyxl')
writer.book = load_workbook(path)
writer.sheets = dict((ws.title,ws) for ws in writer.book.worksheets)

df.to_excel(writer,sheet_name="Winners", startrow=32,startcol = 1,index=False, header=False)
writer.save()
writer.close()

df = pd.DataFrame({'Name':['A','B','C','D'], 'Age': [10,0,30,50]})

writer = pd.ExcelWriter(path, engine='xlsxwriter')

df.to_excel(writer, sheet_name='Winners', startrow=32, startcol=1, index=False)
df.plot(x ='Name', y='Age', kind = 'line')
plt.show()

workbook = writer.book
worksheet = writer.sheets['Winners']
chart = workbook.add_chart({'type': 'column'})

writer.save()
writer.close()

largest_so_far = -1
print('Before', largest_so_far)
for the_num in [9,41,12,3,74,15]:
    if the_num > largest_so_far:
        largest_so_far = the_num
    print(largest_so_far)
print('After', largest_so_far)
